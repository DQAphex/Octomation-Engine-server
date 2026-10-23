import os
from openpyxl import load_workbook,Workbook

class ExcelFile:
  def __init__(
    self,
    file_path: str,
    file_name: str = "gestion_guias_despacho.xlsx"
  ):
    self.file_path = os.path.join(file_path, file_name)

    if os.path.exists(self.file_path):
      self.wb = load_workbook(self.file_path)
    else:
      self.wb = Workbook()

    self.ws = self.wb.active

  def export_to_excel(self, table):
    print(table)
    for row in table:
      self.ws.append(row)

    self.wb.save(self.file_path)